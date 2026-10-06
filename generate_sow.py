from docx import Document

def create_sow():
    doc = Document()
    doc.add_heading('Statement of Work (SOW)', 0)

    doc.add_heading('Project Name', level=1)
    doc.add_paragraph('ITSoftlab360')

    doc.add_heading('Project Overview', level=1)
    doc.add_paragraph('ITSoftlab360 is a comprehensive, modern web application designed to showcase various technological solutions, products, and industry-specific applications. It features public-facing modules such as products, solutions, and franchising, alongside administrative controls and personalized user dashboards with advanced 3D visual elements.')

    doc.add_heading('Business Problem', level=1)
    doc.add_paragraph('There is a need for a centralized, highly interactive, and visually engaging platform to manage and showcase software solutions and products, while also providing secure authentication, role-based dashboards, and a robust backend to handle lead generation and franchise inquiries.')

    doc.add_heading('Proposed Solution', level=1)
    doc.add_paragraph('A full-stack web application built using the latest Next.js (App Router) and React capabilities. The solution utilizes an immersive UI powered by Tailwind CSS, Framer Motion, and Three.js for 3D interactions. The backend is tightly integrated using Next.js Server Components and Drizzle ORM connecting to a PostgreSQL database for reliable data storage and management.')

    doc.add_heading('Product/Modules Overview', level=1)
    
    doc.add_heading('1. Authentication', level=2)
    p_auth = doc.add_paragraph()
    p_auth.add_run('├── Login\n')
    p_auth.add_run('├── Registration\n')
    p_auth.add_run('├── Forgot Password\n')
    p_auth.add_run('└── Reset Password')

    doc.add_heading('2. Public Platform', level=2)
    p_pub = doc.add_paragraph()
    p_pub.add_run('├── Home / Landing Page\n')
    p_pub.add_run('├── Solutions Overview\n')
    p_pub.add_run('├── Products Directory\n')
    p_pub.add_run('├── Industries\n')
    p_pub.add_run('├── Franchise Opportunities\n')
    p_pub.add_run('├── Demo Requests\n')
    p_pub.add_run('└── Contact Us')

    doc.add_heading('3. User Management & Dashboard', level=2)
    p_user = doc.add_paragraph()
    p_user.add_run('├── User Dashboard\n')
    p_user.add_run('└── Profile Management')

    doc.add_heading('4. Administrative Control', level=2)
    p_admin = doc.add_paragraph()
    p_admin.add_run('├── Admin Panel\n')
    p_admin.add_run('├── User & Role Management\n')
    p_admin.add_run('└── Content Management (Products/Solutions)')

    doc.add_heading('Module-wise Functional Requirements', level=1)
    doc.add_paragraph('- Authentication: Secure login, session management, and password recovery workflows.')
    doc.add_paragraph('- Public Pages: SEO-optimized, highly interactive pages showcasing the company’s offerings using 3D elements.')
    doc.add_paragraph('- Franchise: A dedicated funnel for capturing and managing franchise inquiries and leads.')
    doc.add_paragraph('- Dashboard: A secure area for authenticated users to view their interactions and subscriptions.')
    doc.add_paragraph('- Admin: Full CRUD capabilities for managing platform content and users.')

    doc.add_heading('Project Workflow', level=1)
    doc.add_paragraph('1. Requirement Gathering & UI/UX Design (In Progress)\n2. Frontend Development & Component Building (In Progress)\n3. Database Schema Design & ORM Setup (Started)\n4. API Development & Backend Integration (Pending)\n5. Quality Assurance & Testing (Pending)\n6. Deployment & Hosting Setup (Pending)')

    doc.add_heading('Technical Details', level=1)
    doc.add_paragraph('Frontend: React 19, Next.js 16.2.12 (App Router), TypeScript, Tailwind CSS v4, Framer Motion, Three.js / React Three Fiber, React Query')
    doc.add_paragraph('Backend: Next.js API Routes / Server Actions')
    doc.add_paragraph('Database: PostgreSQL, Drizzle ORM')
    doc.add_paragraph('Hosting: To be determined (Recommendation: Vercel / AWS)')
    doc.add_paragraph('APIs: REST / Server Actions for internal communication')
    doc.add_paragraph('Architecture: Serverless Monorepo (Next.js)')
    doc.add_paragraph('Third-Party Integrations: Pending (e.g., Email Service, Payment Gateway)')
    doc.add_paragraph('Dependencies on Other Products/Systems: None identified currently')

    doc.add_heading('Current Status', level=1)
    doc.add_heading('What has been completed?', level=2)
    doc.add_paragraph('- Initial project scaffolding and dependency setup.')
    doc.add_paragraph('- Basic directory structure routing (Home, Admin, Dashboard, Login, Solutions, Products, Franchise).')
    doc.add_paragraph('- Database configuration and initial User schema (Drizzle ORM).')
    
    doc.add_heading('What is currently in progress?', level=2)
    doc.add_paragraph('- Frontend UI development and component creation.')
    doc.add_paragraph('- Integration of 3D elements using Three.js.')
    
    doc.add_heading('What is remaining?', level=2)
    doc.add_paragraph('- Complete backend logic and API implementation.')
    doc.add_paragraph('- Full database schema definition for all modules (Products, Franchise, etc.).')
    doc.add_paragraph('- Authentication flow implementation.')
    doc.add_paragraph('- Admin dashboard functionality and data binding.')
    
    doc.add_heading('Any pending information or dependencies?', level=2)
    doc.add_paragraph('- Confirmation on Hosting environment.')
    doc.add_paragraph('- Required third-party integrations (Payment, Email, CRM).')
    
    doc.add_heading('Any doubts/clarifications required?', level=2)
    doc.add_paragraph('- Detailed functional requirements for the Franchise and Admin modules.')
    doc.add_paragraph('- Specific permission levels needed for User Roles (Admin vs User).')

    doc.add_heading('Weekly Status Update (Template)', level=1)
    doc.add_paragraph('Completed work: [To be filled weekly]')
    doc.add_paragraph('Work in progress: [To be filled weekly]')
    doc.add_paragraph('Remaining work: [To be filled weekly]')
    doc.add_paragraph('Percentage of completion: [X]%')
    doc.add_paragraph('Pending dependencies: [To be filled weekly]')
    doc.add_paragraph('Open questions/clarifications: [To be filled weekly]')

    doc.save('ITSoftlab360_SOW.docx')

if __name__ == "__main__":
    create_sow()
